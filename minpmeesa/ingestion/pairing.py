"""Table d'appariement indicateur / tableau / commentaire (section 2.3.2).

Le programme PROPOSE, l'humain VALIDE. Pour chaque graphique commenté d'un
Rapport d'analyse, on propose le tableau de l'Annuaire du même exercice qui en
porte les valeurs, avec un score combinant :
  - la proximité des intitulés (Jaccard sur des racines de 4 lettres, qui
    neutralisent les coquilles « crées » / « créées ») ;
  - la part des valeurs citées dans le commentaire qui figurent dans le tableau.
Toutes les lignes proposées ont le statut « candidat ». Seul un fichier validé
par double lecture (data/pairing/appariement_valide_AAAA.csv) les fait passer
à « valide » ou « rejete ».

Le code indicateur est stable d'une édition à l'autre : les graphiques des
exercices antérieurs reçoivent le code du graphique le plus récent dont
l'intitulé est le plus proche (au-dessus d'un seuil), sinon un code propre.
"""
from __future__ import annotations

import csv
import re
import unicodedata
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from ..guards.normalize import cles, extraire_nombres

_VIDES = {"les", "des", "de", "du", "la", "le", "en", "et", "par", "selon", "dans", "entre", "au",
          "aux", "d", "l", "a", "sur", "pour", "une", "un", "leur", "leurs", "evolution",
          "repartition", "nombre", "proportion"}


def _plat(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c)).lower().replace("’", "'")


def racines(intitule: str, garder_generiques: bool = True) -> set:
    s = re.sub(r"\([^)]*\)|\b20\d\d\b|\d+", " ", _plat(intitule))
    mots = re.findall(r"[a-z]+", s)
    vides = _VIDES if not garder_generiques else _VIDES - {"evolution", "repartition", "nombre", "proportion"}
    return {m[:4] for m in mots if m not in vides and len(m) > 1}


def jaccard(a: set, b: set) -> float:
    return len(a & b) / len(a | b) if a and b else 0.0


def slug(intitule: str) -> str:
    s = re.sub(r"\([^)]*\)|\b20\d\d\b|\d+", " ", _plat(intitule))
    mots = [m for m in re.findall(r"[a-z]+", s) if m not in _VIDES and len(m) > 2]
    return "-".join(mots[:6]) or "indicateur"


def attribuer_codes(graphiques_par_exercice: Dict[int, list], seuil: float = 0.45) -> Dict[Tuple[int, int], str]:
    """Code stable par (exercice, numéro de graphique). Le plus récent exercice
    fixe les codes ; les antérieurs s'y rattachent un à un par similarité."""
    codes: Dict[Tuple[int, int], str] = {}
    references: Dict[str, set] = {}
    for ex in sorted(graphiques_par_exercice, reverse=True):
        pris = set()
        paires = []
        for g in graphiques_par_exercice[ex]:
            r = racines(g.intitule)
            for code, rr in references.items():
                paires.append((jaccard(r, rr), g.numero, code))
        for s, num, code in sorted(paires, reverse=True):
            if s < seuil or (ex, num) in codes or code in pris:
                continue
            codes[(ex, num)] = code
            pris.add(code)
        for g in graphiques_par_exercice[ex]:
            if (ex, g.numero) not in codes:
                base = slug(g.intitule)
                code, k = base, 2
                while code in references:
                    code, k = f"{base}-{k}", k + 1
                codes[(ex, g.numero)] = code
                references[code] = racines(g.intitule)
    return codes


def score_candidat(graphique, tableau) -> Tuple[float, float, float]:
    """(score, similarité des intitulés, part des valeurs du commentaire
    retrouvées dans le tableau)."""
    s_tit = jaccard(racines(graphique.intitule), racines(tableau.intitule))
    vals = [n for n in extraire_nombres(graphique.commentaire)]
    dans_tab = set()
    for t in tableau.triplets:
        for n in extraire_nombres(t.valeur_texte, bornes_annees=(0, -1), garder_references=True):
            dans_tab |= cles(n)
    s_val = (sum(1 for n in vals if cles(n) & dans_tab) / len(vals)) if vals else 0.0
    return round(0.5 * s_tit + 0.5 * s_val, 3), round(s_tit, 3), round(s_val, 3)


def proposer(graphique, tableaux: list, n: int = 3) -> List[Tuple[float, float, float, object]]:
    """Les n meilleurs tableaux candidats (un tableau scindé sur deux pages
    compte une fois : on fusionne les morceaux de même numéro)."""
    par_num: Dict[int, object] = {}
    for t in tableaux:
        if t.numero is None:
            continue
        if t.numero in par_num:
            par_num[t.numero].triplets = par_num[t.numero].triplets + t.triplets
        else:
            par_num[t.numero] = type(t)(t.page, t.numero, t.intitule, t.bbox, list(t.triplets))
    notes = [(*score_candidat(graphique, t), t) for t in par_num.values()]
    return sorted(notes, key=lambda x: x[0], reverse=True)[:n]


# ---------------------------------------------------------------------------
#  Validation humaine
# ---------------------------------------------------------------------------
CHAMPS_VALIDES = ["code_indicateur", "exercice", "graphique_n", "tableau_n", "statut",
                  "valide_par", "date_validation"]


def charger_validations(dossier: Path) -> Dict[Tuple[str, int], dict]:
    """Lit data/pairing/appariement_valide_AAAA.csv (séparateur « ; »), produit
    après double lecture du fichier a_valider_AAAA.xlsx."""
    out: Dict[Tuple[str, int], dict] = {}
    for p in sorted(dossier.glob("appariement_valide_*.csv")):
        with open(p, encoding="utf-8-sig", newline="") as f:
            for r in csv.DictReader(f, delimiter=";"):
                if r.get("statut") not in ("valide", "rejete"):
                    continue
                out[(r["code_indicateur"], int(r["exercice"]))] = r
    return out


def ecrire_a_valider(lignes: List[dict], chemin: Path) -> None:
    """Fichier Excel de double lecture (une ligne par graphique)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "a_valider"
    entetes = ["code_indicateur", "exercice", "graphique_n", "graphique_intitule",
               "extrait_commentaire", "tableau_n_propose", "tableau_intitule", "page_tableau",
               "extrait_tableau", "score", "sim_intitules", "part_valeurs_retrouvees",
               "autres_candidats", "decision_lecteur_1", "tableau_n_lecteur_1",
               "decision_lecteur_2", "tableau_n_lecteur_2", "statut_final", "tableau_n_final",
               "remarque"]
    ws.append(entetes)
    for l in lignes:
        ws.append([l.get(h, "") for h in entetes])
    gras = Font(bold=True)
    jaune = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = gras
    for col in ("N", "O", "P", "Q", "R", "S", "T"):
        for c in ws[col][1:]:
            c.fill = jaune
    dv = DataValidation(type="list", formula1='"valide,rejete,corrige"', allow_blank=True)
    ws.add_data_validation(dv)
    for col in ("N", "P", "R"):
        dv.add(f"{col}2:{col}{len(lignes) + 1}")
    largeurs = {"A": 34, "D": 50, "E": 70, "G": 50, "I": 60, "M": 40}
    for k, v in largeurs.items():
        ws.column_dimensions[k].width = v
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    notice = wb.create_sheet("mode_emploi")
    for l in [
        "Double lecture de la table d'appariement (mémoire, section 2.3.2).",
        "1. Deux lecteurs remplissent indépendamment les colonnes jaunes « lecteur_1 » et « lecteur_2 » :",
        "   valide = le tableau proposé porte bien les valeurs du graphique ;",
        "   rejete = aucun tableau de l'Annuaire ne correspond ;",
        "   corrige = un autre tableau correspond (indiquer son numéro).",
        "2. Les désaccords sont tranchés ensemble : remplir statut_final (valide/rejete) et tableau_n_final.",
        "3. Lancer : python -m minpmeesa.ingestion.pairing --importer data/pairing/a_valider_AAAA.xlsx --par \"Nom1 / Nom2\"",
        "   Cela écrit data/pairing/appariement_valide_AAAA.csv et calcule le kappa de Cohen entre lecteurs.",
        "4. Relancer la reconstruction : python -m minpmeesa.ingestion.build",
        "Tant que ce fichier n'est pas validé, tous les résultats d'évaluation sont PROVISOIRES.",
    ]:
        notice.append([l])
    notice.column_dimensions["A"].width = 120
    chemin.parent.mkdir(parents=True, exist_ok=True)
    wb.save(chemin)


def importer_validation(xlsx: Path, valide_par: str) -> Path:
    """Convertit un fichier a_valider_AAAA.xlsx rempli en appariement_valide_AAAA.csv
    et affiche l'accord entre lecteurs (kappa de Cohen sur la décision)."""
    import datetime

    from openpyxl import load_workbook
    ws = load_workbook(xlsx, read_only=True)["a_valider"]
    rows = list(ws.iter_rows(values_only=True))
    h = {k: i for i, k in enumerate(rows[0])}
    out, d1, d2 = [], [], []
    for r in rows[1:]:
        if r[h["decision_lecteur_1"]] and r[h["decision_lecteur_2"]]:
            d1.append(str(r[h["decision_lecteur_1"]]))
            d2.append(str(r[h["decision_lecteur_2"]]))
        statut = (r[h["statut_final"]] or "").strip()
        if statut not in ("valide", "rejete"):
            continue
        tab = r[h["tableau_n_final"]] or r[h["tableau_n_propose"]]
        out.append({"code_indicateur": r[h["code_indicateur"]], "exercice": r[h["exercice"]],
                    "graphique_n": r[h["graphique_n"]], "tableau_n": tab, "statut": statut,
                    "valide_par": valide_par, "date_validation": datetime.date.today().isoformat()})
    ex = re.search(r"(20\d\d)", xlsx.name).group(1)
    dest = xlsx.parent / f"appariement_valide_{ex}.csv"
    with open(dest, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=CHAMPS_VALIDES, delimiter=";", extrasaction="ignore")
        w.writeheader()
        w.writerows(out)
    if d1:
        from ..evaluation.metrics import kappa_cohen
        print(f"Kappa de Cohen entre lecteurs ({len(d1)} lignes) : {kappa_cohen(d1, d2):.3f}")
    print(f"{len(out)} lignes validées ou rejetées -> {dest}")
    return dest


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="Import d'un fichier d'appariement validé")
    ap.add_argument("--importer", required=True)
    ap.add_argument("--par", required=True, help="noms des lecteurs")
    a = ap.parse_args()
    importer_validation(Path(a.importer), a.par)
