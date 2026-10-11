"""Export Word (python-docx) des commentaires et des notes."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from ..generation.strategic_note import RUBRIQUES

GRIS = RGBColor(0x59, 0x59, 0x59)
ROUGE = RGBColor(0xA0, 0x1E, 0x1E)


def _rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa.lstrip("#").upper())


def _doc(titre: str) -> Document:
    from ..app import identite
    i = identite.charger()
    vert = _rgb(i["couleurs"].get("vert_fonce", i["couleurs"]["vert"]))
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    for niveau in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        d.styles[niveau].font.color.rgb = vert
    # En-tête de page : logo (s'il est fourni) + nom de l'outil et de la structure
    ent = d.sections[0].header.paragraphs[0]
    lg = identite.logo(i)
    if lg:
        ent.add_run().add_picture(str(lg), height=Cm(1.2))
        ent.add_run("   ")
    r = ent.add_run(f"{i['nom']} — {i['structure']}")
    r.font.size, r.font.color.rgb = Pt(9), GRIS
    d.add_heading(titre, level=0)
    return d


def _avertissements(d: Document, obj: dict):
    msgs = []
    if obj.get("mode_extractif"):
        msgs.append("texte rédigé par gabarits, à partir des seuls chiffres de l'Annuaire (sans modèle de langage).")
    if obj.get("provisoire"):
        msgs.append("provisoire : les liens entre graphiques et tableaux n'ont pas encore été validés par la Cellule.")
    if obj.get("commentaires_non_valides") or obj.get("valide") is False:
        msgs.append("commentaires non encore validés par la Cellule : document à relire.")
    for m in msgs:
        p = d.add_paragraph()
        r = p.add_run("Avertissement : " + m)
        r.bold, r.font.color.rgb, r.font.size = True, ROUGE, Pt(9)


def _note_source(d: Document, texte: str):
    p = d.add_paragraph()
    r = p.add_run(texte)
    r.italic, r.font.size, r.font.color.rgb = True, Pt(8), GRIS
    p.paragraph_format.space_after = Pt(2)


def _commentaire(d: Document, res: dict, niveau: int = 1):
    if res.get("abstention"):
        p = d.add_paragraph()
        r = p.add_run(f"{res.get('mention', 'sources insuffisantes, à commenter manuellement')} — {res.get('motif', '')}")
        r.italic = True
        return
    for e in res.get("enonces", []):
        p = d.add_paragraph(e["texte"], style="List Bullet" if e.get("type") == "perspective" else None)
        from ..guards.provenance import a_afficher
        refs = "; ".join(a_afficher(e.get("references", [])))
        if refs:
            _note_source(d, "Sources : " + refs)
    if res.get("valeurs_ecartees"):
        _note_source(d, "Valeurs écartées par le contrôle : " +
                     ", ".join(f"{v['valeur']} ({v['action']})" for v in res["valeurs_ecartees"]))


def exporter_commentaire(res: dict, chemin: Path) -> Path:
    d = _doc(f"Commentaire — {res.get('indicateur') or res.get('tableau_intitule')}")
    from ..guards.provenance import nom_document
    _note_source(d, f"Exercice {res['exercice']} ; {nom_document(res.get('doc_annuaire'))}, tableau "
                    f"{res.get('tableau_n')} ; rédaction : {res.get('modele')}")
    _avertissements(d, res)
    if not res.get("abstention"):
        from ..generation.redaction import structurer_commentaire
        c = structurer_commentaire(res)
        d.add_heading("Commentaire", level=1)
        _para(d, c["paragraphe"])
        if c["source"]:
            _note_source(d, c["source"])
        d.add_heading("Détail des sources, énoncé par énoncé", level=2)
    _commentaire(d, res)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    d.save(chemin)
    return chemin


def _para(d: Document, texte: str, style: str | None = None):
    p = d.add_paragraph(texte, style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p


def exporter_note_analyse(note: dict, chemin: Path) -> Path:
    """Rapport : introduction, principaux enseignements, chapitres > sections > tableaux (un paragraphe
    rédigé et sa source par tableau), annexes. Le détail énoncé par énoncé est dans l'interface."""
    from ..generation.redaction import structurer_note_analyse
    doc = structurer_note_analyse(note)
    d = _doc(note["titre"])
    _note_source(d, f"Établie le {note['date']} — rédaction : {note['modele']}")
    _avertissements(d, note)
    d.add_heading("Introduction", level=1)
    _para(d, doc["introduction"])
    if doc["enseignements"]:
        d.add_heading("Principaux enseignements", level=1)
        for x in doc["enseignements"]:
            _para(d, x, style="List Bullet")
    for c in doc["chapitres"]:
        d.add_heading(c["titre"], level=1)
        for se in c["sections"]:
            if se["titre"]:
                d.add_heading(se["titre"], level=2)
            for t in se["tableaux"]:
                d.add_heading(t["titre"], level=3)
                if t.get("transition"):
                    _para(d, t["transition"])
                _para(d, t["paragraphe"])
                if t["source"]:
                    _note_source(d, t["source"] + ("" if t.get("valide") else " Commentaire non encore validé."))
    if doc["manuels"] or doc["incoherences"]:
        d.add_heading("Annexes", level=1)
    if doc["manuels"]:
        d.add_heading("A. Tableaux à commenter par la Cellule", level=2)
        for a in doc["manuels"]:
            d.add_paragraph(f"Tableau {a['tableau_n']} — {a.get('intitule') or a['code']}", style="List Bullet")
    if doc["incoherences"]:
        d.add_heading("B. Contrôle de cohérence de l'Annuaire", level=2)
        for inc in doc["incoherences"]:
            det = " ; ".join(f"{ref} (tableau(x) {', '.join(map(str, n))})" for ref, n in inc["references"].items())
            _para(d, f"Pour un même total {inc['exercice']} ({inc['total']}), la valeur {inc['exercice'] - 1} "
                     f"diffère selon le tableau : {det}. Les évolutions calculées diffèrent donc ; à vérifier.",
                  style="List Bullet")
    chemin.parent.mkdir(parents=True, exist_ok=True)
    d.save(chemin)
    return chemin


def exporter_note_strategique(note: dict, chemin: Path) -> Path:
    d = _doc(note["titre"])
    _note_source(d, f"Établie le {note['date']} ; {note['nb_evolutions']} évolutions retenues ; "
                    f"{note['nb_mots']} mots ; modèle : {note['modele']}")
    _avertissements(d, note)
    if note.get("chiffres_cles"):
        from ..app import identite
        i = identite.charger()
        t = d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        cel = t.rows[0].cells[0]
        titre = cel.paragraphs[0].add_run("Chiffres clés")
        titre.bold, titre.font.color.rgb = True, _rgb(i["couleurs"].get("vert_fonce", i["couleurs"]["vert"]))
        for b in note["chiffres_cles"]:
            cel.add_paragraph(b["texte"], style="List Bullet")
        d.add_paragraph()
    for r in RUBRIQUES:
        d.add_heading(r, level=1)
        for b in note["rubriques"][r]:
            d.add_paragraph(b["texte"], style="List Bullet" if r != "Évolutions marquantes" else None)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Chaque valeur est tracée jusqu'au tableau et à la page (voir le fichier de traçabilité).")
    r.italic, r.font.size = True, Pt(8)
    chemin.parent.mkdir(parents=True, exist_ok=True)
    d.save(chemin)
    return chemin
