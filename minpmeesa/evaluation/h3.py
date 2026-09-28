"""H3 : temps de réponse du système, fichiers de la séance de chronométrage et
grille BN1 à BN5 (Annexe IV). Les temps humains et les notes des lecteurs ne
sont JAMAIS remplis par le programme : ils sont saisis par l'étudiant, puis
`run_all --h3` calcule médianes, réduction et conformité.
"""
from __future__ import annotations

import statistics as st
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

CRITERES = {
    "BN1": "Sélection : les 5 à 7 évolutions retenues sont les plus marquantes",
    "BN2": "Mise en perspective : tendance ou variation ponctuelle correctement qualifiée",
    "BN3": "Rattachement : objectifs de politique publique cités seulement s'ils sont documentés",
    "BN4": "Traçabilité : chaque énoncé renvoie au commentaire, au tableau et à la page",
    "BN5": "Format : 1 à 2 pages, rubriques attendues, pistes sans chiffre nouveau",
}
ENTETE = PatternFill("solid", fgColor="1F4E79")


def _entete(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = ENTETE


def protocole_docx(chemin: Path, taches: list[dict]) -> Path:
    d = Document()
    d.add_heading("Protocole de chronométrage (H3, section 2.4.4)", 0)
    for t in [
        "Objectif : mesurer le temps nécessaire à un cadre de la Cellule pour produire le commentaire d'un "
        "tableau de l'Annuaire, avec et sans le système.",
        "Participants : au moins 3 cadres. Chaque cadre traite deux tableaux DIFFÉRENTS : l'un avec le système, "
        "l'autre sans ; l'ordre des deux conditions est alterné d'un cadre à l'autre (voir le plan ci-dessous).",
        "Sans le système : Annuaire, rapports antérieurs et tableur habituels. Avec le système : parcours "
        "« Commenter un indicateur », relecture, corrections, validation.",
        "Le chronomètre démarre à la remise du tableau et s'arrête quand le cadre déclare son commentaire prêt "
        "à être publié (relecture comprise).",
        "Saisir les heures de début et de fin dans h3_saisie_temps.xlsx, sans arrondir ; noter toute interruption "
        "dans la colonne remarques. Ne pas communiquer les temps aux participants pendant la séance.",
        "Calcul : python -m minpmeesa.evaluation.run_all --h3 data/results/<date>/h3_saisie_temps.xlsx",
    ]:
        d.add_paragraph(t, style="List Number")
    d.add_heading("Plan de la séance", 1)
    tb = d.add_table(rows=1, cols=4)
    tb.style = "Light Grid Accent 1"
    for i, h in enumerate(["Cadre", "Ordre", "Condition", "Tableau / indicateur"]):
        tb.rows[0].cells[i].text = h
    for t in taches:
        r = tb.add_row().cells
        r[0].text, r[1].text, r[2].text, r[3].text = t["cadre"], str(t["ordre"]), t["condition"], t["indicateur"]
    d.save(chemin)
    return chemin


def plan(indicateurs: list[dict], n_cadres: int = 3) -> list[dict]:
    """Deux tableaux différents par cadre, ordre des conditions alterné."""
    out = []
    for k in range(n_cadres):
        a, b = indicateurs[(2 * k) % len(indicateurs)], indicateurs[(2 * k + 1) % len(indicateurs)]
        conds = ["avec", "sans"] if k % 2 == 0 else ["sans", "avec"]
        for o, (c, ind) in enumerate(zip(conds, (a, b)), start=1):
            out.append({"cadre": f"Cadre {k + 1}", "ordre": o, "condition": c,
                        "indicateur": f"Tableau {ind['tableau_n']} — {ind['tableau_intitule']}"})
    return out


def saisie_xlsx(chemin: Path, taches: list[dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "saisie"
    _entete(ws, ["cadre", "ordre", "condition", "indicateur", "heure_debut (hh:mm:ss)", "heure_fin (hh:mm:ss)",
                 "duree_min", "remarques"])
    for i, t in enumerate(taches, start=2):
        ws.append([t["cadre"], t["ordre"], t["condition"], t["indicateur"], None, None,
                   f'=IF(AND(E{i}<>"",F{i}<>""),(TIMEVALUE(F{i})-TIMEVALUE(E{i}))*1440,"")', ""])
    for col, w in zip("ABCDEFGH", [10, 7, 10, 60, 20, 20, 10, 30]):
        ws.column_dimensions[col].width = w
    wb.save(chemin)
    return chemin


def grille_xlsx(chemin: Path, notes: list[dict]) -> Path:
    """Une ligne par (note, critère) ; BN4 et BN5 pré-remplis par la vérification automatique."""
    wb = Workbook()
    ws = wb.active
    ws.title = "grille"
    _entete(ws, ["note", "fichier", "critere", "libelle", "verification_automatique",
                 "lecteur_1 (0/1/2)", "lecteur_2 (0/1/2)", "remarques"])
    for n in notes:
        v = n["verification"]
        auto = {"BN4": "satisfait" if v["BN4_toute_valeur_tracee"] else "NON satisfait",
                "BN5": ("satisfait" if v["BN5_longueur_ok"] and v["BN5_rubriques"] and v["BN5_messages_cles_3"]
                        else "NON satisfait") + f" ({v['nb_mots']} mots)"}
        for c, lib in CRITERES.items():
            ws.append([n["titre"], n["fichier"], c, lib, auto.get(c, "non vérifiable automatiquement"), None, None, ""])
    for col, w in zip("ABCDEFGH", [40, 32, 7, 70, 32, 16, 16, 30]):
        ws.column_dimensions[col].width = w
    wb.save(chemin)
    return chemin


def calculer_temps(chemin: Path) -> dict:
    """Réduction du temps médian à partir de la saisie remplie par l'étudiant."""
    from datetime import datetime, time as dtime
    ws = load_workbook(chemin, data_only=False)["saisie"]
    avec, sans = [], []

    def minutes(x):
        if x is None or x == "":
            return None
        if isinstance(x, dtime):
            return x.hour * 60 + x.minute + x.second / 60
        if isinstance(x, datetime):
            return x.hour * 60 + x.minute + x.second / 60
        h = [int(p) for p in str(x).split(":")] + [0, 0]
        return h[0] * 60 + h[1] + h[2] / 60

    for r in ws.iter_rows(min_row=2, values_only=True):
        d, f = minutes(r[4]), minutes(r[5])
        if d is None or f is None:
            continue
        (avec if r[2] == "avec" else sans).append(f - d)
    if not avec or not sans:
        return {"statut": "non mesuré", "raison": "temps non encore saisis"}
    ma, ms = st.median(avec), st.median(sans)
    return {"statut": "mesuré", "n_avec": len(avec), "n_sans": len(sans), "mediane_avec_min": round(ma, 2),
            "mediane_sans_min": round(ms, 2), "reduction": round(1 - ma / ms, 4) if ms else None,
            "n_cadres": len({r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[4]})}


def calculer_grille(chemin: Path, seuil_moy: float, crit_min: int) -> dict:
    ws = load_workbook(chemin)["grille"]
    notes: dict[str, dict] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[5] is None or r[6] is None:
            continue
        notes.setdefault(r[0], {})[r[2]] = (float(r[5]) + float(r[6])) / 2
    if not notes:
        return {"statut": "non mesuré", "raison": "notes des deux lecteurs non encore saisies"}
    detail = {n: {"moyennes": c, "criteres_satisfaits": sum(v >= seuil_moy for v in c.values())}
              for n, c in notes.items()}
    conformes = sum(d["criteres_satisfaits"] >= crit_min for d in detail.values())
    return {"statut": "mesuré", "n_notes": len(detail), "notes_conformes": conformes,
            "part_conformes": round(conformes / len(detail), 4), "detail": detail}
