"""Étape 7 : calculs H3 à partir des fichiers remplis, verdicts, métriques."""
from openpyxl import load_workbook

from minpmeesa.evaluation import h3, metrics as M, run_all


def test_h3_temps_et_grille(tmp_path):
    taches = h3.plan([{"tableau_n": i, "tableau_intitule": f"T{i}"} for i in range(1, 7)])
    assert len(taches) == 6 and {t["condition"] for t in taches} == {"avec", "sans"}
    assert taches[0]["condition"] != taches[2]["condition"]           # ordre alterné
    p = h3.saisie_xlsx(tmp_path / "s.xlsx", taches)
    assert h3.calculer_temps(p)["statut"] == "non mesuré"            # rien n'est inventé
    wb = load_workbook(p)
    ws = wb["saisie"]
    for r in range(2, 8):
        ws[f"E{r}"] = "10:00:00"
        ws[f"F{r}"] = "10:20:00" if ws[f"C{r}"].value == "sans" else "10:08:00"
    wb.save(p)
    t = h3.calculer_temps(p)
    assert t["mediane_sans_min"] == 20 and t["mediane_avec_min"] == 8 and t["reduction"] == 0.6
    notes = [{"titre": "N", "fichier": "n.docx", "verification": {"BN4_toute_valeur_tracee": True,
              "BN5_longueur_ok": True, "BN5_rubriques": True, "BN5_messages_cles_3": True, "nb_mots": 600}}]
    g = h3.grille_xlsx(tmp_path / "g.xlsx", notes)
    assert h3.calculer_grille(g, 1.5, 4)["statut"] == "non mesuré"
    wb = load_workbook(g)
    ws = wb["grille"]
    for r, (a, b) in zip(range(2, 7), [(2, 2), (2, 1), (1, 1), (2, 2), (2, 1)]):
        ws[f"F{r}"], ws[f"G{r}"] = a, b
    wb.save(g)
    r = h3.calculer_grille(g, 1.5, 4)
    assert r["detail"]["N"]["criteres_satisfaits"] == 4 and r["part_conformes"] == 1.0


def test_verdicts():
    assert run_all.verdict([("a", True), ("b", True)]) == "validée"
    assert run_all.verdict([("a", True), ("b", None)]) == "non concluante"
    assert run_all.verdict([("a", False), ("b", None)]) == "non validée"


def test_metriques_recuperation():
    assert M.succes_a_k([False, True], 1) == 0 and M.succes_a_k([False, True], 5) == 1
    assert M.rr([False, False, True]) == 1 / 3
    assert M.rouge_l("les PME créées en 2024", "les PME créées en 2024") == 1.0
    assert M.est_pertinent({"doc_id": "d", "page": 3, "page_fin": 5}, [{"doc_id": "d", "pages": [4]}])


def test_choix_encodeur_compare_au_repli(base_complete, tmp_path, monkeypatch):
    """Un candidat présent est mesuré et comparé au repli, question par question."""
    from minpmeesa.evaluation import choix_encodeur
    from minpmeesa.evaluation.h1_recuperation import charger_jeu
    from minpmeesa.retrieval.hybrid import Moteur
    from minpmeesa.store import encodeur as encmod

    class Faux(encmod.EncodeurLSA):          # se comporte comme un encodeur neuronal présent sur le disque
        def __init__(self, nom, *args, **kwargs):
            super().__init__(64, 42)
            self.nom, self._ajuste = nom, False

        def encoder_passages(self, textes):
            self.ajuster(textes)
            return super().encoder_passages(textes)

    cible = "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
    monkeypatch.setattr(encmod, "EncodeurST", Faux)
    monkeypatch.setattr(encmod, "dossier_modele", lambda n: tmp_path if n == cible else tmp_path / "absent")
    res = choix_encodeur.comparer(Moteur(), charger_jeu(), tmp_path, [cible])
    assert [e["encodeur"] for e in res] == [encmod.NOM_REPLI, cible]
    assert all(e["statut"] == "mesuré" for e in res)
    w = res[1]["wilcoxon_mrr_hybride_vs_repli"]
    assert w["n"] == 52 and 0 <= res[1]["hybride_succes_5"] <= 1
    lignes = (tmp_path / "tableau_3_3_par_question.csv").read_text(encoding="utf-8-sig").splitlines()
    assert len(lignes) == 53 and f"{cible}|hybride|rr" in lignes[0]


def test_choix_encodeur_reprise(base_complete, tmp_path):
    """Un encodeur déjà mesuré est repris de partiel/, à l'identique, sans être recalculé."""
    from minpmeesa.evaluation import choix_encodeur
    from minpmeesa.evaluation.h1_recuperation import charger_jeu
    from minpmeesa.retrieval.hybrid import Moteur
    m, jeu = Moteur(), charger_jeu()
    r1 = choix_encodeur.comparer(m, jeu, tmp_path, [])
    assert list((tmp_path / "partiel").glob("*.json"))
    r2 = choix_encodeur.comparer(m, jeu, tmp_path, [])
    assert r1 == r2
