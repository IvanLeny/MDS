"""Étape 5 : note d'analyse, note stratégique, exports."""
import pytest

from minpmeesa.export import docx_export
from minpmeesa.generation import analysis_note, commentary, strategic_note
from minpmeesa.guards import literal_check as lc
from minpmeesa.retrieval.hybrid import Moteur


@pytest.fixture(scope="module")
def m(base_complete):
    return Moteur()


def test_note_analyse_assemble_sans_valeur_nouvelle(base_complete, m):
    note = analysis_note.rediger(base_complete, 2024, chapitre=r"CHAPITRE I\b", moteur=m, journaliser=False)
    assert note["sections"]
    nums = [s["tableau_n"] for s in note["sections"]]
    assert nums == sorted(nums)                           # ordre du plan de l'Annuaire
    for s in note["sections"]:
        seul = commentary.commenter(base_complete, s["code"], 2024, None, journaliser=False, moteur=m)
        assert [e["texte"] for e in s["commentaire"]["enonces"]] == [e["texte"] for e in seul["enonces"]]


def test_note_analyse_prefere_les_commentaires_valides(base_complete, m):
    code = "pme-creees-sur-periode"
    res = commentary.commenter(base_complete, code, 2024, None, journaliser=True, moteur=m)
    analysis_note.valider_production(base_complete, res["prod_id"], "test")
    v = analysis_note.commentaire_valide(base_complete, code, 2024)
    assert v and v["prod_id"] == res["prod_id"] and v["valide"]
    note = analysis_note.rediger(base_complete, 2024, codes=[code], moteur=m, journaliser=False)
    assert note["sections"][0]["valide"] and not note["commentaires_non_valides"]
    analysis_note.valider_production(base_complete, res["prod_id"], "test", "a_relire")


def test_note_strategique_bn(base_complete, m):
    note = strategic_note.rediger(base_complete, 2024, None, m, journaliser=False)
    assert 1 <= note["nb_evolutions"] <= 7
    assert list(note["rubriques"]) == strategic_note.RUBRIQUES
    v = strategic_note.verifier_bn4_bn5(note)
    assert v["BN4_toute_valeur_tracee"] and v["BN5_longueur_ok"] and v["BN5_messages_cles_3"]
    # pistes pour la décision : aucun chiffre
    assert all(not lc.nombres(b["texte"]) for b in note["rubriques"]["Pistes pour la décision"])
    # chaque variation retenue est tracée à ses deux valeurs sources
    for c in note["selection"]:
        assert sum(1 for t in note["tracabilite"] if t["code_indicateur"] == c["code"] and t["variation"]) == 2
    # un objectif n'est rattaché que s'il est documenté (avec page)
    for c in note["selection"]:
        assert c["objectif"] is None or (c["objectif"]["page"] and c["objectif"]["doc_id"])


def test_qualification_tendance():
    assert "tendance à la hausse" in strategic_note.qualifier([1.0, 2.0, 3.0], 3)
    assert "ponctuelle" in strategic_note.qualifier([1.0, -2.0, 3.0], 3)
    assert "trop court" in strategic_note.qualifier([1.0], 3)


def test_exports_word(base_complete, m, tmp_path):
    from docx import Document
    res = commentary.commenter(base_complete, "pme-creees-sur-periode", 2024, None, journaliser=False)
    p = docx_export.exporter_commentaire(res, tmp_path / "c.docx")
    texte = "\n".join(x.text for x in Document(p).paragraphs)
    assert "MODE EXTRACTIF" in texte and "21 132" in texte
    note = strategic_note.rediger(base_complete, 2024, None, m, journaliser=False)
    p = docx_export.exporter_note_strategique(note, tmp_path / "n.docx")
    t = "\n".join(x.text for x in Document(p).paragraphs)
    assert all(r in t for r in strategic_note.RUBRIQUES)


class _FauxLLM:
    nom, moteur, modele = "faux:llm", "faux", "llm"

    def __init__(self, reponse):
        self.reponse = reponse

    def generer(self, systeme, utilisateur, json_attendu=True):
        return self.reponse


@pytest.fixture
def commentaires_valides(base_complete, m):
    """La note stratégique part des commentaires VALIDÉS : on valide ceux produits en mode
    gabarits, puis on rétablit l'état initial."""
    ids = []
    for (code,) in base_complete.execute("SELECT code_indicateur FROM appariement WHERE exercice=2024 "
                                         "AND statut IN ('valide','candidat') AND tableau_n IS NOT NULL"):
        r = commentary.commenter(base_complete, code, 2024, None, journaliser=True, moteur=m)
        if not r.get("abstention"):
            analysis_note.valider_production(base_complete, r["prod_id"], "test")
            ids.append(r["prod_id"])
    yield ids
    for i in ids:
        analysis_note.valider_production(base_complete, i, "test", "a_relire")


def test_note_strategique_redigee_par_le_llm_et_controlee(base_complete, m, commentaires_valides):
    import json
    ref = strategic_note.rediger(base_complete, 2024, None, m, journaliser=False)
    c0 = ref["selection"][0]["code"]
    rep = json.dumps({
        "messages_cles": ["Le stock progresse nettement.", "Les créations restent dynamiques.",
                          "Les UPA reculent de 99,9 %."],                       # chiffre inventé
        "evolutions": [{"code": c0, "mise_en_perspective": "Cette hausse s'inscrit dans une tendance durable."}],
        "points_attention": ["Le recul des UPA appelle une vigilance."],
        "pistes": ["Renforcer l'accompagnement des artisans.", "Viser 25 000 créations par an."]},  # chiffre interdit
        ensure_ascii=False)
    # le programme sélectionne toujours (même sélection) : seul le client change
    note = strategic_note.rediger(base_complete, 2024, _FauxLLM(rep), m, journaliser=False)
    assert [c["code"] for c in note["selection"]] == [c["code"] for c in ref["selection"]]   # BN1 : le programme
    assert note["redaction"] == "modèle de langage" and note["moteur"] == "faux"
    textes = " ".join(b["texte"] for r in strategic_note.RUBRIQUES for b in note["rubriques"][r])
    assert "99,9" not in textes and "25 000" not in textes
    assert {e["valeur"] for e in note["valeurs_ecartees"]} >= {"99,9", "25 000"}
    assert "tendance durable" in textes


def test_note_strategique_repli_gabarits_si_json_invalide(base_complete, m, commentaires_valides):
    note = strategic_note.rediger(base_complete, 2024, _FauxLLM("pas du json"), m, journaliser=False)
    assert note["redaction"].startswith("gabarits (repli")
    assert strategic_note.verifier_bn4_bn5(note)["BN4_toute_valeur_tracee"]


def test_reformulation_consultation_controlee(m):
    import json
    from minpmeesa.retrieval import consultation
    f = _FauxLLM(json.dumps({"reponse": "Les CFCE ont enregistré 21 132 PME en 2024. Cela représente 55 555 emplois."}))
    r = consultation.consulter("Combien de PME ont été créées dans les CFCE en 2024 ?", m,
                               journaliser=False, rediger=True, client=f)
    assert "21 132" in r["reponse"]["texte"] and "55 555" not in r["reponse"]["texte"]
    assert r["reponse"]["valeurs_ecartees"][0]["valeur"] == "55 555"


def test_lot_b_note_strategique(base_complete):
    """B1 : messages clés de débuts différents, chacun chiffré ; B2 : >= 3 types de pistes ;
    B3 : encadré chiffres clés ; note 2023 : >= 4 évolutions ; contrôles BN4/BN5 toujours satisfaits."""
    from minpmeesa.generation import strategic_note as sn
    from minpmeesa.guards import literal_check as lc
    n = sn.rediger(base_complete, 2024, journaliser=False)
    msgs = [b["texte"] for b in n["rubriques"]["Messages clés"]]
    assert len(msgs) == 3 and len({m.split()[0] for m in msgs}) == 3
    assert all(lc.nombres(m) for m in msgs)
    types = {b["texte"].split(" :")[0] for b in n["rubriques"]["Pistes pour la décision"]}
    assert len(types) >= 3
    assert n["chiffres_cles"] and all(b.get("sources") for b in n["chiffres_cles"])
    v = sn.verifier_bn4_bn5(n)
    assert v["BN4_toute_valeur_tracee"] and v["BN5_longueur_ok"]
    assert sn.rediger(base_complete, 2023, journaliser=False)["nb_evolutions"] >= 4


def test_de_contraction():
    from minpmeesa.generation.strategic_note import de
    assert de("le stock de PME") == "du stock de PME"
    assert de("les UPA enregistrées") == "des UPA enregistrées"
    assert de("la valeur ajoutée des PME") == "de la valeur ajoutée des PME"
