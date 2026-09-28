"""Étape 2 : ingestion du corpus réel (valeurs connues, détection, appariement)."""
import pymupdf

from minpmeesa import config
from minpmeesa.ingestion import registry, tables


def _texte(nom):
    d = pymupdf.open(str(config.chemin("corpus") / nom))
    return "\n".join(p.get_text() for p in d)


def test_exercice_par_le_contenu_malgre_titre_trompeur():
    # L'Annuaire 2023 porte le titre courant « Annuaire statistique 2022 ».
    t = _texte("annuaire_2023.pdf")
    assert "ANNUAIRE STATISTIQUE 2022" in t
    assert registry.detecter(t) == {"type": "annuaire", "exercice": 2023, "trimestre": None}


def test_trimestre_note_conjoncture():
    assert registry.detecter(_texte("note_conjoncture_T4_2023.pdf"))["trimestre"] == "T4"


def test_valeurs_tableau_1_annuaire_2024(base_complete):
    r = base_complete.execute(
        "SELECT valeur_texte, annee_colonne FROM valeurs WHERE doc_id='annuaire_2024' AND tableau_n=1 "
        "AND ligne='PME' AND sous_colonne != '%' ORDER BY annee_colonne").fetchall()
    serie = {a: v for v, a in r if a}
    assert serie[2024] == "443 524" and serie[2023] == "393 175" and serie[2016] == "202 746"


def test_tableau_libelle_sur_rangee_distincte(base_complete):
    # Tableau 2 : les libellés de secteur sont sous leurs valeurs dans la grille PDF.
    r = base_complete.execute(
        "SELECT valeur_texte FROM valeurs WHERE doc_id='annuaire_2024' AND tableau_n=2 "
        "AND ligne='Tertiaire' AND annee_colonne=2024 AND sous_colonne='%'").fetchone()
    assert r[0] == "77,2"


def test_libelle_contenant_un_nombre():
    mots = [(10, 10, 40, 20, "Moins"), (42, 10, 52, 20, "de"), (54, 10, 64, 20, "20"),
            (66, 10, 80, 20, "ans"), (150, 10, 170, 20, "3"), (172, 10, 190, 20, "341"),
            (150, 0, 170, 8, "2016"),
            (10, 30, 40, 40, "Total"), (150, 30, 170, 40, "202"), (172, 30, 190, 40, "746")]
    trip = tables.lineariser_mots(mots)
    assert [(t.ligne, t.valeur_texte, t.annee_colonne) for t in trip] == [
        ("Moins de 20 ans", "3 341", 2016), ("Total", "202 746", 2016)]


def test_valeur_num_ecritures_francaises():
    assert tables.valeur_num("12 500") == 12500
    assert tables.valeur_num("7,8") == 7.8
    assert tables.valeur_num("79,6 %") == 79.6
    assert tables.valeur_num("(3,2)") == -3.2


def test_appariement_2024_corrige_l_erreur_connue(base_complete):
    # L'ancienne table appariait le graphique 2 (stock par secteur) au tableau 17 (VA).
    r = base_complete.execute("SELECT tableau_n, statut FROM appariement WHERE exercice=2024 AND graphique_n=2").fetchone()
    assert r[0] == 2


def test_codes_indicateurs_stables_entre_editions(base_complete):
    codes = {r[0] for r in base_complete.execute(
        "SELECT code_indicateur FROM appariement GROUP BY code_indicateur HAVING COUNT(DISTINCT exercice) >= 2")}
    assert "stock-pme-region" in codes


def test_fichiers_de_validation_produits():
    d = config.chemin("appariement")
    for ex in (2021, 2022, 2023, 2024):
        assert (d / f"a_valider_{ex}.xlsx").exists()
        assert (d / f"appariement_{ex}.csv").exists()
    assert (config.chemin("corpus") / "inventaire.csv").exists()


def test_import_double_lecture_et_kappa(tmp_path):
    from minpmeesa.ingestion import pairing
    from openpyxl import load_workbook
    lignes = [pairing.Ligne(f"code{i}", 2099, i, f"G{i}", i, f"T{i}", "r", "a", 0.9, 0.9, 0.9)
              for i in range(1, 5)]
    pairing.ecrire_csv(lignes, tmp_path / "appariement_2099.csv")
    x = pairing.ecrire_xlsx_validation(lignes, tmp_path / "a_valider_2099.xlsx")
    wb = load_workbook(x)
    ws = wb["a_valider"]
    decisions = [("valide", "valide"), ("valide", "rejete"), ("rejete", "rejete"), ("valide", "valide")]
    for k, (a, b) in enumerate(decisions, start=2):
        ws[f"O{k}"], ws[f"Q{k}"] = a, b
    wb.save(x)
    bilan = pairing.importer_validation(x, tmp_path)
    assert bilan["valide"] == 2 and bilan["rejete"] == 1 and bilan["a_verifier"] == 1
    assert bilan["kappa_cohen"] is not None
    statuts = [r["statut"] for r in pairing.lire_csv(tmp_path / "appariement_2099.csv")]
    assert statuts == ["valide", "a_verifier", "rejete", "valide"]
